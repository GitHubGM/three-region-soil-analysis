"""Build the three-domain ML manuscript from verified experiment tables."""
from __future__ import annotations
import csv
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/multiregion'
QA=ROOT/'analysis/multiregion_manuscript'
QA.mkdir(parents=True,exist_ok=True)
SOURCE=ROOT/'Hybrid_Early_Warning_Soil_Degradation_Kazakhstan_EcolInform.docx'
DEST=ROOT/'Machine_Learning_Soil_Moisture_Three_Domains.docx'
TITLE='Spatial validation and regional transfer of machine learning models for soil moisture reconstruction'
LABELS={'Aral':'Aral Sea','Atyrau':'Atyrau','Bozaigyr':'Northwestern'}
NAMES={'training_mean':'Training mean','ridge_env':'Ridge E','svr_env':'SVR E','rf_geo':'RF G',
       'rf_env':'RF E','rf_env_geo':'RF E + G','gbr_env':'GBR E','gbr_env_geo':'GBR E + G'}
PROTOCOLS=['random_rows','random_sites','buffered_quadrants','leave_region_out']
summary=json.loads((OUT/'summary.json').read_text(encoding='utf-8'))
literature=json.loads((ROOT/'analysis/multiregion_literature.json').read_text(encoding='utf-8'))

def read(name):
    with (OUT/name).open(newline='',encoding='utf-8') as f:return list(csv.DictReader(f))

region_rows=read('main_region_metrics.csv')
M={(r['protocol'],r['region'],r['model']):r for r in region_rows}
macro_rows=read('main_macro_metrics.csv')
MAC={(r['protocol'],r['model']):r for r in macro_rows}
A={(r['region'],r['model'],float(r['fraction']),r['setting']):r for r in read('adaptation_region_metrics.csv')}
AM={(r['model'],float(r['fraction']),r['setting']):r for r in read('adaptation_macro_metrics.csv')}
SHIFT={(r['region'],r['feature']):r for r in read('feature_shift.csv')}
SENS={(r['model'],float(r['fraction']),r['setting']):r for r in read('sensitivity/macro_seed_ranges.csv')}
best={p:min((r for r in macro_rows if r['protocol']==p),key=lambda r:float(r['rmse'])) for p in PROTOCOLS}
hash_before=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
doc=Document(SOURCE)
authors=[x.text for x in doc.paragraphs[1:4] if x.text.strip()]
for child in list(doc._element.body):
    if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption','List Bullet']:
    style=doc.styles[name]
    style.font.name='Times New Roman'
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):del fonts.attrib[key]
    for key in ['ascii','hAnsi','cs','eastAsia']:fonts.set(qn('w:'+key),'Times New Roman')
    style.font.color.rgb=RGBColor(0,0,0)
    color=style.element.rPr.find(qn('w:color'))
    if color is not None:
        for key in list(color.attrib):
            if key!=qn('w:val'):del color.attrib[key]
    if style.element.pPr is not None:
        for tag in ['w:pBdr','w:shd']:
            node=style.element.pPr.find(qn(tag))
            if node is not None:style.element.pPr.remove(node)
normal=doc.styles['Normal']
normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.space_after=Pt(9)
normal.paragraph_format.widow_control=True
normal.paragraph_format.keep_with_next=False
normal.paragraph_format.keep_together=False
for level,size,before in [(1,14,20),(2,12,10),(3,11,8)]:
    s=doc.styles[f'Heading {level}']
    s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(5)
    s.paragraph_format.keep_with_next=True;s.paragraph_format.keep_together=True
doc.styles['Title'].font.size=Pt(15)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].paragraph_format.line_spacing=1
doc.styles['Caption'].paragraph_format.space_after=Pt(8)

def p(text,style='Normal'):
    a=doc.add_paragraph(text,style)
    if style=='Normal':a.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    previous=a._p.getprevious()
    if previous is not None and previous.tag==qn('w:tbl'):a.paragraph_format.space_before=Pt(6)
    return a

def h(text,level=1):return p(text,f'Heading {level}')

def table(caption,headers,rows,widths,numeric_cols=()):
    cap=p(caption,'Caption');cap.paragraph_format.keep_with_next=True
    for run in cap.runs:run.bold=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    for cells,values,is_header in [(t.rows[0].cells,headers,True)]+[(t.add_row().cells,r,False) for r in rows]:
        for i,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=str(value)
            tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                node=OxmlElement('w:'+side);node.set(qn('w:w'),'75' if side in ['top','bottom'] else '85');node.set(qn('w:type'),'dxa');mar.append(node)
            tcpr.append(mar)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EAF0' if is_header else 'FFFFFF');tcpr.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1
                par.paragraph_format.keep_with_next=False
                par.alignment=WD_ALIGN_PARAGRAPH.CENTER if i in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
                for run in par.runs:run.font.name='Times New Roman';run.font.size=Pt(9.5);run.bold=is_header
        pr=cells[0]._tc.getparent().get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if is_header:pr.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        node=OxmlElement('w:'+side);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
    t._tbl.tblPr.append(borders)
    return t

def fig(filename,caption,width=6):
    a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER
    a.paragraph_format.space_before=Pt(2);a.paragraph_format.space_after=Pt(0)
    a.paragraph_format.line_spacing=1;a.paragraph_format.keep_with_next=True
    shape=a.add_run().add_picture(str(OUT/'figures'/filename),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    a=p(caption,'Caption');a.paragraph_format.keep_together=True

def f(value,digits=4):return f'{float(value):.{digits}f}'

title=p(TITLE,'Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
for author in authors:
    a=p(author);a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_after=Pt(4)
    for r in a.runs:r.font.size=Pt(10)
h('Highlights')
for text in [
    'Eight regression configurations are compared under four sampling and transfer protocols.',
    'Shared locations and neighbouring source cells are excluded from regional-transfer training.',
    'Environmental and geographic predictors are assessed on identical test locations.',
    'Target supervision is evaluated with nested samples of spatially separated labelled sites.'
]:p(text,'List Bullet')
h('Abstract')
transfer_name=best['leave_region_out']['model']
transfer_values=[float(M['leave_region_out',r,transfer_name]['rmse']) for r in LABELS]
p(f'Performance estimates for geospatial machine learning depend on how training and evaluation locations are separated. We study this dependence using ERA5-family reanalysis soil moisture across the Aral Sea, Atyrau and a Northwestern domain. Twenty-four NetCDF files provide eight variables on 10 and 18 August 2026, with 7,534 distinct physical sites. Two extraction domains share 336 locations. We compare eight fixed regression configurations under random record splits, random site splits, buffered spatial blocks and regional holdouts. Regional transfer excludes the entire target rectangle and a 0.3° coordinate buffer from source data. Models use environmental predictors, geographic coordinates or both; scaling and weighting use training records only. The lowest arithmetic mean of regional root mean squared errors is {f(best["random_sites"]["rmse"])} m³/m³ under random site sampling, {f(best["buffered_quadrants"]["rmse"])} m³/m³ under buffered blocks and {f(best["leave_region_out"]["rmse"])} m³/m³ under regional transfer. The best transfer configuration, gradient boosting with environmental predictors, has regional errors from {min(transfer_values):.4f} to {max(transfer_values):.4f} m³/m³. Further experiments compare target-only learning with source-plus-target learning at three labelled-site budgets. The study provides a reproducible assessment of validation design, predictor choice and supervised regional transfer for contemporaneous reconstruction of modelled soil water. The results apply to the two archived dates and the specified extraction domains.')
p('Keywords: geospatial machine learning; regional generalization; spatial cross-validation; distribution shift; supervised transfer; soil moisture').alignment=WD_ALIGN_PARAGRAPH.LEFT

h('1 Introduction')
p('Machine-learning models fitted to gridded environmental data are often evaluated on records drawn from the same spatial distribution as their training set. This can be a useful estimate of interpolation accuracy, but deployment may require predictions for a separated area or an unobserved region. Repeated dates at the same location create another distinction: a randomly withheld record can have a second observation from that location in training. A meaningful evaluation must state which of these settings it represents and retain that distinction when comparing models (Roberts et al., 2017; Ploton et al., 2020).')
p('Regional soil-moisture reconstruction provides a tractable regression setting for studying these questions. Reanalysis products contain spatially aligned soil-water, temperature, wind and vegetation fields. Their common grid allows models to be evaluated on the same targets under alternative splits and predictor sets. At the same time, neighbouring cells share spatial structure and regional distributions differ. A model that performs well when test cells are interspersed with training cells may therefore behave differently when applied to an entire held-out domain. This connects an environmental application to the machine-learning problem of distribution shift (Koh et al., 2021).')
p('The present study uses three rectangular extraction domains in and around Kazakhstan. The Aral Sea and Atyrau domains provide contrasting geographic settings, while a smaller Northwestern domain partly overlaps the northern portion of the Atyrau extraction. This overlap makes physical-site identity an explicit experimental concern. Treating folder names as independent domains would allow observations at the same coordinates to enter both source and target sets. We identify shared locations from their coordinates and exclude them, together with a surrounding coordinate buffer, before fitting regional-transfer models.')
p('We address three questions. First, how do error estimates and model rankings change between random records, grouped sites, spatial blocks and regional holdouts? Second, how does adding geographic position to environmental predictors affect reconstruction within and across domains? Third, when labelled target-region sites become available, does combining them with source-region data improve prediction over target-only training? The contribution is an empirical comparison with inspectable split memberships, predictions and independent verification. The analysis concerns reconstruction of reanalysis soil moisture at the recorded dates; its prediction target is a modelled environmental state.')

h('2 Related work')
for i,text in enumerate(literature['draft_paragraphs']):
    text=text.replace('Aral, Atyrau or Bozaigyr','the Aral Sea, Atyrau or Northwestern domain')
    if i==2:
        text=text.replace('an explicit regional transfer task.','an explicit regional transfer task after shared sites and nearby source cells have been excluded.')
    if i==5:
        text=('Building on these studies, we compare validation protocols on identical regional targets and examine the effect of geographic predictors. A second experiment varies the amount of target-region supervision while keeping test blocks fixed. Using standard regressors makes differences in sampling, source-domain composition and feature choice easier to inspect. Domain-specific error tables and complete predictions support comparisons beyond a single model ranking.')
    p(text)

h('3 Data and prediction task')
h('3.1 Extraction domains and observations',2)
p('Each domain contains eight NetCDF variables on a regular 0.1° latitude–longitude grid. The two observation dates are 10 and 18 August 2026. Coordinates align exactly among the eight files within each folder, and all 125,920 variable values are finite. The archives contain 15,740 region-labelled site-date records. Because the Atyrau and Northwestern domains intersect, these represent 15,068 distinct physical site-date records at 7,534 locations. Table 1 and Figure 1 define the footprints; they are rectangular extraction domains and do not follow administrative, land-cover or shoreline boundaries.')
table('Table 1  Regional extraction domains', ['Domain','Latitude °N','Longitude °E','Sites','Records'],[
    [LABELS[r['folder']],f'{r["lat_min"]:.1f}–{r["lat_max"]:.1f}',f'{r["lon_min"]:.1f}–{r["lon_max"]:.1f}',f'{r["n_sites"]:,}',f'{r["n_rows"]:,}'] for r in summary['regions']
],[1.5,1.15,1.15,1.1,1.1],numeric_cols=(1,2,3,4))
p('The Northwestern label refers to 48.5–52.1°N and 50.6–53.3°E. The corresponding archive folder is named Bozaigyr; this name is retained only as a provenance identifier in the code and machine-readable results. Study geography is determined from the NetCDF coordinates. The shared Atyrau–Northwestern footprint contains 336 physical locations and 672 site-date records. Shared soil-water, wind and leaf-area values agree exactly. On 18 August, the two soil-temperature fields differ between archives by at most 0.0000611 K, so duplicate source records are resolved using fixed alphabetical folder priority.')
fig('figure_1_domains.png','Figure 1. Coordinate footprints of the three extraction domains. The overlap between the Atyrau and Northwestern domains contains 336 grid locations. Rectangles show data coverage; their edges are not administrative or water boundaries.')
h('3.2 Target and predictors',2)
p('The target is volumetric soil water in the 0–7 cm model layer, swvl1, expressed in m³/m³. The environmental feature set E contains soil temperatures in the 0–7 cm and 7–28 cm layers, the signed eastward wind component at 10 m and low-vegetation leaf area index. Geographic features G are latitude and longitude. Soil temperatures are converted from kelvin to degrees Celsius. The wind component remains a signed velocity. Leaf area index is interpreted as prescribed seasonal vegetation structure under the ERA5-Land formulation (Muñoz-Sabater et al., 2021). Date and region identifier are excluded from model inputs.')
p('Snow cover, lake-bottom temperature and lake-ice thickness are retained as archive diagnostics. Snow cover is zero throughout the data, and lake-ice values are numerical roundoff near zero. Values in swvl1 and lake-ice thickness between −10⁻¹² and zero are set to zero in the analysis copy. The raw inputs are preserved. Regional averages describe the full extraction rectangles; lake diagnostic fields are not used to infer a water mask.')
p('The files identify ECMWF as the source institution and record daily-mean aggregation with a five-hour time shift. We use ERA5-family terminology because the archive metadata are the basis of product attribution. The daily-processing framework allows configurable sampling and time shifts (ECMWF, 2026). Input hashes fix the exact archived release used in this experiment. Model scores measure agreement with these reanalysis values, with the temporal scope defined by the two dates.')

h('4 Experimental methods')
h('4.1 Regression configurations',2)
p('Eight configurations separate algorithm choice from feature-set choice (Table 2). The training-mean baseline predicts the weighted mean target in each training split. Ridge regression and support vector regression (SVR) with a radial basis function (RBF) kernel provide linear and kernel comparators. Random forests (RF) and gradient-boosting regression (GBR) provide tree-ensemble comparators (Hoerl and Kennard, 1970; Smola and Schölkopf, 2004; Breiman, 2001; Friedman, 2001). A coordinate-only forest evaluates how much predictive information is associated with location. The environmental and combined feature sets are compared separately for both tree ensembles.')
table('Table 2  Fixed model configurations',['Configuration','Set','Settings'],[
    ['Training mean','None','Weighted training mean'],['Ridge E','E','Weighted standardization; α = 10'],
    ['SVR E','E','Weighted standardization; RBF; C = 1; ε = 0.005; γ = scale'],
    ['RF G','G','200 trees; minimum leaf 10; all features per split'],
    ['RF E','E','Same forest settings'],['RF E + G','E + G','Same forest settings'],
    ['GBR E','E','150 stages; learning rate 0.03; depth 2; minimum leaf 10'],
    ['GBR E + G','E + G','Same boosting settings'],
],[1.25,.6,4.15])
p('The primary model settings, protocols and label budgets were fixed before the three-domain scores were calculated. No held-out score is used for parameter tuning or feature selection. The model random seed and primary sampling-seed base are 20261005. Forests use bootstrap sampling and squared-error splits; boosting uses squared-error loss. Standardization for ridge and SVR is fitted on training data with the same sample weights used by the regressor. Predictions are retained without clipping. Thus, the comparison evaluates these explicit configurations rather than optimized representatives of each algorithm family.')
h('4.2 Sampling and regional transfer protocols',2)
p('Random-record evaluation uses four shuffled folds within each region. It deliberately permits the other date at a held-out location to appear in training and serves as a conventional record-splitting diagnostic. Random-site evaluation instead partitions the unique coordinate pairs into four shuffled groups and assigns both dates at a site to the same fold. Nearby sites can still be distributed between training and test sets. Both designs train only on records from the same regional domain as their test data.')
p('Buffered spatial evaluation divides each regional rectangle into four quadrants at the midpoints of its latitude and longitude extents. A quadrant is withheld in turn, and training sites within a 0.3° rectangular coordinate buffer around it are removed. Both dates remain together. For a candidate point, the exclusion distance is the maximum of its nonnegative latitude and longitude distances to the held-out rectangle. A point is retained only when that distance exceeds 0.3° plus a numerical tolerance of 10⁻⁸. The buffer is specified in coordinate degrees and is not a constant ground distance.')
p('Regional holdout uses one complete domain as the target and the other two archives as sources. Before fitting, source sites within the complete target rectangle or its 0.3° buffer are excluded. This removes both coordinate duplicates and nearby observations across overlapping archive boundaries. Any remaining source-source duplicates are reduced to one record per rounded coordinate pair and date, using alphabetical folder priority. Coordinates are rounded to six decimal places when physical-site IDs are constructed. These rules operate on coordinates rather than folder membership.')
h('4.3 Weighting and evaluation',2)
p('Within a single region, cell weights are proportional to the cosine of latitude, accounting for cell-area changes on the regular angular grid. In regional transfer, weights are normalized so each retained source-region provenance contributes equal total weight after deduplication. All training weights are then scaled to have mean one. This prevents the larger retained source from determining the objective solely through its cell count. Evaluation weights within each target region are proportional to cosine latitude.')
p('We calculate root mean squared error (RMSE), mean absolute error (MAE), signed bias and R². Bias is prediction minus target. For each model and protocol, all held-out predictions from a region are combined before calculating its metrics. Regional R² uses the weighted mean of that region’s evaluation targets as the reference. Macro RMSE is the arithmetic mean of the three regional RMSEs; it is not a pooled cell-level RMSE. Relative MSE skill is one minus the ratio of a model’s MSE to that of the corresponding training-mean baseline. Complete fold scores are retained to expose geographic variation.')
h('4.4 Target supervision experiment',2)
p('A second experiment evaluates environmental random forests and gradient boosting with three labelled-target budgets. In each buffered target quadrant split, the eligible training sites are placed in a seeded random order. Nested subsets contain 5%, 20% or 100% of those eligible sites, rounded upward to a whole site. Both dates at a selected site are included. The percentage refers to the eligible training pool after the held-out block and its buffer have been removed. Identical subsets are used for both models and both training settings.')
p('Target-only training uses those labelled sites. Source-plus-target training adds the two source archives after the same complete-target-rectangle purge used in regional holdout. Consequently, source files cannot reveal labels at unselected target sites. The target subset receives half the total training weight and each retained source provenance receives one quarter; cosine-latitude weighting operates within each group. This is pooled supervised learning with target labels. The source-only regional model provides a zero-label reference on the same evaluation locations. Each regional budget score pools predictions across its four test quadrants.')
p('After the primary curves showed nonmonotonic changes with label budget, a post hoc sensitivity analysis repeated the 5% and 20% selections with subset-seed bases 20271005 and 20281005. Each base is offset by 100 times the region index plus the fold index. Model seeds, source pools, test blocks, weighting and all hyperparameters remain fixed. The complete 100% subset is unchanged and is not retrained. We report the range across the primary and two additional subset seeds as descriptive sampling sensitivity, not a confidence interval; no seed is chosen according to its score.')
h('4.5 Predictor support and reproducibility',2)
p('To describe transfer conditions, each environmental predictor’s minimum and maximum are computed in the purged source pool. We report the area-weighted fraction of target cell-date records outside each range and outside at least one range. With two equally represented dates, this is the mean of the two daily area fractions. This univariate envelope diagnostic does not implement the multivariate area-of-applicability method of Meyer and Pebesma (2021). Scores inside and outside the envelope are also exported as descriptive error strata. Geographic feature comparisons use identical folds and targets; they do not identify causal environmental effects.')
p(f'The implementation uses Python 3.12.14 and scikit-learn {summary["packages"]["scikit-learn"]} (Pedregosa et al., 2011), with all package versions and fixed settings archived alongside the outputs. The main comparison contains {summary["main_fit_count"]} configuration evaluations and {summary["main_prediction_count"]:,} predictions. The primary target-supervision experiment adds {summary["adaptation_fit_count"]} fits and {summary["adaptation_prediction_count"]:,} predictions. The post hoc subset sensitivity adds 192 fits and 251,840 predictions. Preprocessing, split memberships, model errors and file hashes are independently checked. Runtime is recorded per main configuration and split. Repeated cell observations, spatial dependence and the partial domain overlap are retained in the interpretation; no cell-based significance test is used to claim generality across regions.')

h('5 Results')
h('5.1 Regional contrasts and source retention',2)
p('The regional soil-water means differ substantially. On 10 August, area-weighted mean swvl1 was 0.0589 m³/m³ in the Aral Sea domain, 0.1018 m³/m³ in Atyrau and 0.1596 m³/m³ in the Northwestern domain. The corresponding values on 18 August were 0.0607, 0.1027 and 0.1646 m³/m³. The Northwestern domain had the smallest daily spatial standard deviation, approximately 0.0422–0.0435 m³/m³. These contrasts provide context for both regional bias and R², whose reference variance is specific to each target domain.')
p('After overlap and buffer removal, the regional-transfer training pools contained 5,341 sites for the Aral Sea target, 2,809 for Atyrau and 6,324 for the Northwestern target. For Aral Sea transfer, source-source deduplication removed 672 repeated site-date records. For Atyrau and Northwestern transfer, exclusion of the target footprint and buffer removed 840 and 1,020 source records, respectively. Every retained training pool was disjoint from its target at the physical-site level.')
h('5.2 Validation changes error estimates and rankings',2)
rf_site=float(MAC['random_sites','rf_env_geo']['rmse'])
rf_block=float(MAC['buffered_quadrants','rf_env_geo']['rmse'])
rf_transfer=float(MAC['leave_region_out','rf_env_geo']['rmse'])
p(f'The combined random forest had the lowest macro RMSE under random-record and random-site evaluation, {f(MAC["random_rows","rf_env_geo"]["rmse"])} and {rf_site:.4f} m³/m³, respectively. Its macro error rose to {rf_block:.4f} m³/m³ under buffered quadrants and {rf_transfer:.4f} m³/m³ under regional holdout. Relative to random sites, these errors were {100*(rf_block/rf_site-1):.1f}% and {100*(rf_transfer/rf_site-1):.1f}% higher. The protocols differ in training composition and spatial support as well as separation, so these percentages describe the evaluated prediction settings.')
p(f'Environmental gradient boosting had the lowest macro RMSE under buffered blocks, {f(MAC["buffered_quadrants","gbr_env"]["rmse"])} m³/m³, and regional holdout, {f(MAC["leave_region_out","gbr_env"]["rmse"])} m³/m³. Environmental ridge regression ranked second in regional transfer at {f(MAC["leave_region_out","ridge_env"]["rmse"])} m³/m³, ahead of the environmental forest at {f(MAC["leave_region_out","rf_env"]["rmse"])} m³/m³. Thus, the configuration leading random-site prediction did not lead the separated-domain comparison (Figure 2).')
fig('figure_2_macro_rmse.png','Figure 2. Macro RMSE for eight fixed configurations under four evaluation protocols. Each entry is the arithmetic mean of three regional RMSEs. Values are in m³/m³; the protocols represent different spatial prediction settings.')
h('5.3 Geographic features and regional transfer',2)
rf_env_site=float(MAC['random_sites','rf_env']['rmse'])
rf_env_transfer=float(MAC['leave_region_out','rf_env']['rmse'])
gbr_env_transfer=float(MAC['leave_region_out','gbr_env']['rmse'])
gbr_geo_transfer=float(MAC['leave_region_out','gbr_env_geo']['rmse'])
p(f'Adding coordinates reduced random-site forest RMSE by {100*(1-rf_site/rf_env_site):.1f}%, but increased its regional-transfer RMSE by {100*(rf_transfer/rf_env_transfer-1):.1f}%. For gradient boosting, adding coordinates increased regional-transfer macro RMSE from {gbr_env_transfer:.4f} to {gbr_geo_transfer:.4f} m³/m³, an {100*(gbr_geo_transfer/gbr_env_transfer-1):.1f}% increase. The coordinate-only forest was competitive under random-site sampling at {f(MAC["random_sites","rf_geo"]["rmse"])} m³/m³ but reached {f(MAC["buffered_quadrants","rf_geo"]["rmse"])} m³/m³ under buffered blocks. Geographic predictors therefore supplied information that was useful under interpolation-like sampling but less transferable in these tests.')
table('Table 3  Regional holdout performance of environmental gradient boosting',
      ['Target domain','RMSE','MAE','Bias','R²'],
      [[LABELS[r]]+[f(M['leave_region_out',r,'gbr_env'][key],3 if key=='r2' else 4) for key in ['rmse','mae','bias','r2']] for r in LABELS],
      [1.8,1.05,1.05,1.05,1.05],numeric_cols=(1,2,3,4))
p('Error and bias units are m³/m³. Each target domain uses its own area-weighted mean as the R² reference. This table reports the same configuration across domains; it does not select a different model for each domain.','Caption')
p('Regional rankings differed. Combined gradient boosting gave the smallest Aral Sea transfer RMSE, 0.0405 m³/m³. Environmental random forest led Atyrau at 0.0579 m³/m³. Environmental gradient boosting led the Northwestern target at 0.0459 m³/m³, but its R² was −0.143. Every tested configuration had negative Northwestern transfer R² (Figure 3). The same domain also had negative R² for all fitted regressors under buffered quadrants, where the training-mean baseline had the lowest error. Positive performance in the other domains therefore does not establish useful spatial reconstruction throughout the study area.')
fig('figure_3_region_transfer.png','Figure 3. Domain-specific R² under regional holdout. Each value uses the target domain’s weighted target mean as its reference. Negative values indicate errors greater than that reference; the reference is calculated for scoring and is not a deployable source-only predictor.')
h('5.4 Target supervision and source pooling',2)
benefits=sum(float(A[r,m,frac,'source_plus_target']['rmse'])<float(A[r,m,frac,'target_only']['rmse'])
             for r in LABELS for m in ['rf_env','gbr_env'] for frac in [.05,.2,1.0])
p(f'Source-plus-target pooling produced lower regional RMSE than target-only learning in {benefits} of the 18 matched domain–model–budget comparisons in the primary experiment. Figure 4 retains all comparisons and shows the source-only reference. A comparison uses identical test cells and exactly the same labelled target sites in both settings. Differences therefore reflect the added source records and the specified training weights, conditional on the fixed models and selected sites.')
for model in ['rf_env','gbr_env']:
    zero=MAC['leave_region_out',model]['rmse']
    vals=[AM[model,frac,'source_plus_target']['rmse'] for frac in [.05,.2,1.0]]
    alone=[AM[model,frac,'target_only']['rmse'] for frac in [.05,.2,1.0]]
    p(f'For {NAMES[model]}, source-plus-target macro RMSE at 5%, 20% and 100% of eligible labelled sites was {f(vals[0])}, {f(vals[1])} and {f(vals[2])} m³/m³. Target-only errors at the same budgets were {f(alone[0])}, {f(alone[1])} and {f(alone[2])} m³/m³. The zero-label regional-transfer reference was {f(zero)} m³/m³. These values describe one nested sample per fold, with both dates retained at every selected site.')
fig('figure_4_target_adaptation.png','Figure 4. Effect of labelled target sites on environmental random forest and gradient boosting. Budgets are percentages of eligible target training sites after block and buffer exclusion. Lines compare target-only and source-plus-target training; horizontal references use zero target labels. Test sites are identical across budgets and training settings. Panels retain separate regional scales.')
sens_rows=[]
for model in ['rf_env','gbr_env']:
    for fraction in [.05,.2]:
        values=[]
        for setting in ['target_only','source_plus_target']:
            a=SENS[model,fraction,setting]
            values.append(f'{float(a["rmse_min"]):.4f}–{float(a["rmse_max"]):.4f}')
        sens_rows.append([NAMES[model],f'{fraction:.0%}']+values)
table('Table 4  Macro RMSE ranges across three target subset seeds',
      ['Configuration','Target sites','Target only','Source plus target'],sens_rows,
      [1.3,.85,1.925,1.925],numeric_cols=(1,2,3))
p('Ranges are minimum to maximum macro RMSE in m³/m³ across the primary subset seed and two post hoc repetitions. They describe variation with site selection and are not confidence intervals. Full-budget subsets are identical across seeds.','Caption')
p('The repeated selections show that the primary curves are sensitive to which target sites are labelled. At 5%, target-only random forest macro RMSE ranged from 0.0459 to 0.0491 m³/m³, below the full-budget value of 0.0511 m³/m³ in all three runs. For target-only gradient boosting, the 5% range was 0.0468–0.0498 m³/m³ and straddled the full-budget value of 0.0484 m³/m³. The source-plus-target gradient-boosting range at 20% was narrower, 0.0461–0.0464 m³/m³. These repetitions preserve the primary scores and provide a sampling check on their interpretation.')
h('5.5 Environmental support and computational record',2)
outside=[100*float(SHIFT[r,'any_environment_feature']['target_fraction_outside']) for r in LABELS]
p(f'The area-weighted fractions of target cell-date records outside at least one source predictor range were {outside[0]:.2f}% for the Aral Sea domain, {outside[1]:.2f}% for Atyrau and {outside[2]:.2f}% for the Northwestern domain. Most records were therefore inside every separate source range. This did not guarantee strong transfer: all Northwestern R² scores remained negative. The test examines marginal ranges and does not establish coverage of joint predictor combinations or invariance of the target relationship.')
runtime=read('runtime.csv')
fit_means={m:sum(float(r['fit_seconds']) for r in runtime if r['model']==m)/39 for m in NAMES}
p(f'Mean preprocessing-and-fit times across the 39 main splits were {fit_means["ridge_env"]:.2f} s for ridge regression, {fit_means["svr_env"]:.2f} s for SVR, {fit_means["rf_env"]:.2f} s for environmental random forest and {fit_means["gbr_env"]:.2f} s for environmental gradient boosting. The forests used two worker threads. These wall-clock measurements describe this implementation and its unequal split sizes, rather than a hardware-independent speed ranking. Package versions, system metadata and individual timing records are included in the experiment archive.')

h('6 Discussion')
h('6.1 Implications for geospatial model evaluation',2)
p('The main computational finding is a change in the leading configuration when the evaluation task changes. A combined forest leads randomly interspersed site prediction, while environmental gradient boosting leads the average error under separated blocks and regional transfer. The strong random-site coordinate baseline and the loss of its advantage under blocking are consistent with the spatial-validation concerns raised by Ploton et al. (2020). The practical implication is that a model selected for interpolation should be reassessed before regional deployment. Random-site grouping protects repeated locations but still permits information shared by nearby training and test cells.')
p('There is no monotonic difficulty ordering that every model must follow. Environmental gradient boosting gives similar macro errors under spatial blocks and regional transfer, although the source composition and retained training sizes differ. Individual domains also behave differently. These outcomes support reporting the complete split rules and regional errors, rather than describing the most separated protocol as universally hardest. The interpretation follows Roberts et al. (2017): validation should represent the intended prediction setting, including the distribution of available training data.')
h('6.2 Overlap control and feature choice',2)
p('Archive membership alone is an insufficient unit for regional separation. The two western extractions share nearly one third of the Northwestern sites. Excluding physical duplicates and nearby source cells makes the regional experiment interpretable as prediction at unseen locations. It does not turn the three extraction tasks into independent geographic replicates. Target sites in the overlap are evaluated in two distinct regional tasks, and macro scores summarize those specified tasks. This dependence is why the analysis reports descriptive comparisons rather than significance claims based on thousands of nominally independent cells.')
p('Coordinate features improve interpolation-like performance but increase average transfer error for both ensemble families. This is a feature-set result under fixed configurations, not a claim that coordinates are always harmful. The Aral Sea target provides a counterexample: combined boosting slightly improves its regional RMSE. The broader result is that spatial position can provide local predictive information whose usefulness changes across domains. Marginal predictor ranges supply complementary information, but the Northwestern failure shows that staying inside those ranges is insufficient to establish transferability (Meyer and Pebesma, 2021).')
h('6.3 Learning with target labels',2)
p('The target-supervision experiment changes what information is available to the learner while preserving each spatial test block. Nested site selection and matched comparisons make source pooling directly comparable with target-only training at each budget. Source-only performance is a useful zero-label reference, but adding sources is not assumed to help. The direction and size of its effect must be read by domain, model and budget. Because the target records are reanalysis labels, the experiment quantifies the value of access to target-domain model states rather than the benefit of purchasing field measurements.')
p('The source-plus-target weighting gives target labels a fixed half-share regardless of budget and represents one explicit learning policy. The three subset seeds reveal sampling sensitivity, but label budget also changes the effective constraint imposed by the fixed minimum leaf size. For example, a 5% Northwestern target-only fit contains 68 records, compared with 1,324–1,358 records at the full budget, while the minimum leaf remains 10. Smaller training sets therefore impose a larger relative leaf-size constraint. The curves assess this complete fixed-configuration procedure. Repeated site selections do not isolate that regularization effect, and the small-budget advantage is not evidence that discarding labels generally improves prediction.')
h('6.4 Scope of the evidence',2)
p('The study compares established learning methods on a small number of regional domains. Its computational contribution lies in the controlled validation, physical overlap handling, feature comparisons and supervised transfer evaluation. The fixed configurations support an inspectable benchmark but do not establish optimal hyperparameters or a universal ordering of algorithm families. Regional R² and the Northwestern baseline results make the practical limits of the fitted relationships explicit.')
p('Both dates are included in spatial training pools, so the results concern contemporaneous reconstruction. The two snapshots do not estimate seasonal or long-term performance. Predictors and target originate from the same reanalysis system, and agreement measures internal reconstruction of that system’s fields. Full extraction rectangles include all archived grid locations. These scope conditions define the scientific interpretation of the scores without changing the central finding that sampling design and regional transfer alter the apparent usefulness of model complexity and geographic features.')

h('7 Conclusions')
p(f'Eight fixed regression configurations were evaluated across the Aral Sea, Atyrau and Northwestern extraction domains using 7,534 distinct sites observed on two August 2026 dates. Shared locations were detected by physical coordinates, and regional-transfer training excluded the complete target rectangle and its buffer. The combined random forest achieved macro RMSE {rf_site:.4f} m³/m³ under random-site evaluation, but {rf_transfer:.4f} m³/m³ under regional transfer. Environmental gradient boosting had the lowest macro error under buffered blocks and regional holdout, {f(MAC["buffered_quadrants","gbr_env"]["rmse"])} and {gbr_env_transfer:.4f} m³/m³, respectively.')
p(f'Adding coordinates increased transfer error on average, despite helping randomly interspersed-site prediction. All Northwestern transfer R² values were negative, and source-plus-target training improved over target-only training in {benefits} of 18 matched budget comparisons in the primary experiment. These results support reporting deployment-relevant spatial tests, domain-specific errors and explicit target-supervision conditions when assessing geospatial regression. The archived split memberships and predictions make the findings reproducible for the specified data and configurations.')

h('Declaration of competing interest')
p('The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.')
h('Data and code availability')
p('The reproducibility archive contains the 24 input-file hashes, decoded raw and analysis tables, the settings fixed before scoring, complete train–test memberships and weights, and all held-out predictions. Python scripts generate the experiments, independent checks, figures and manuscript. The underlying reanalysis products are distributed through the Copernicus Climate Data Store. Code and derived tables accompany the manuscript in the reproducibility archive.')
h('References')
use={'roberts2017','ploton2020','koh2021','meyer2021','munoz2021','c3s2024daily','ecmwf_daily_docs'}
refs=[r for r in literature['references'] if r['id'] in use]
refs+=json.loads((ROOT/'analysis/method_references.json').read_text(encoding='utf-8'))['references']
for ref in sorted(refs,key=lambda r:r['authors']):
    text=f'{ref["authors"]} ({ref["year"]}). {ref["title"]}.'
    if ref.get('venue'):text+=' '+ref['venue']
    if ref.get('volume'):text+=', '+ref['volume']
    if ref.get('issue'):text+='('+ref['issue']+')'
    if ref.get('pages'):text+=', '+ref['pages']
    elif ref.get('article'):text+=', '+ref['article']
    text=text.rstrip('.')+'. '+('https://doi.org/'+ref['doi'] if ref.get('doi') else ref['url'])
    if ref.get('accessed'):text+=' (accessed 5 October 2026).'
    a=p(text);a.alignment=WD_ALIGN_PARAGRAPH.LEFT
    a.paragraph_format.space_after=Pt(7);a.paragraph_format.keep_together=True
    a.paragraph_format.left_indent=Inches(.2);a.paragraph_format.first_line_indent=Inches(-.2)
    for run in a.runs:run.font.size=Pt(10)

verification=json.loads((OUT/'verification.json').read_text(encoding='utf-8'))
assert verification['failed']==0 and verification['passed']>0
sensitivity_verification=json.loads((OUT/'sensitivity/verification.json').read_text(encoding='utf-8'))
assert sensitivity_verification['failed']==0 and sensitivity_verification['passed']>0
doc.core_properties.title=TITLE
doc.core_properties.subject='Spatial machine learning and supervised regional transfer'
doc.core_properties.keywords='machine learning, spatial validation, regional transfer, soil moisture'
doc.core_properties.comments=''
doc.save(DEST)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==hash_before
reopened=Document(DEST)
text='\n\n'.join(x.text for x in reopened.paragraphs)
(QA/'manuscript_text.txt').write_text(text,encoding='utf-8')
manifest={'source':str(SOURCE),'source_sha256':hash_before,'output':str(DEST),
          'output_sha256':hashlib.sha256(DEST.read_bytes()).hexdigest(),
          'paragraphs':len(reopened.paragraphs),'tables':len(reopened.tables),'figures':len(reopened.inline_shapes),
          'verification_passed':verification['passed'],'verification_failed':verification['failed'],
          'sensitivity_verification_passed':sensitivity_verification['passed'],
          'title':TITLE,'northwestern_folder':'Bozaigyr','input_data_unchanged':summary['all_inputs_unchanged']}
(QA/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps(manifest,indent=2))

