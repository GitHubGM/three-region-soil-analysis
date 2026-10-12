"""Build an Aral-only revision from the source manuscript and verified results."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Hybrid_Early_Warning_Soil_Degradation_Kazakhstan_EcolInform.docx'
OUTPUT = ROOT / 'Aral_Sea_Soil_Moisture_Revised.docx'
QA = ROOT / 'analysis' / 'manuscript_revision'
QA.mkdir(parents=True, exist_ok=True)
summary = json.loads((ROOT/'results/experiment_summary.json').read_text(encoding='utf-8'))
review = json.loads((ROOT/'analysis/aral_literature_review.json').read_text(encoding='utf-8'))
metrics = {(x['design'],x['model']):x for x in summary['model_metrics']}
changes = {x['variable']:x for x in summary['change_summary']}
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
doc = Document(SOURCE)
source_author_lines = [p.text for p in doc.paragraphs[1:4]]

# This revision replaces the scientific body while retaining source page geometry,
# styles and author metadata. The original file remains the retained source.
for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'):
        doc._element.body.remove(child)

for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption','List Bullet']:
    s=doc.styles[name]
    s.font.name='Times New Roman'
    fonts=s.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        del fonts.attrib[key]
    for key in ['ascii','hAnsi','cs','eastAsia']:
        fonts.set(qn('w:'+key),'Times New Roman')
    s.font.color.rgb=RGBColor(0,0,0)
    if s.element.rPr is not None:
        color=s.element.rPr.find(qn('w:color'))
        if color is not None:
            for key in list(color.attrib):
                if key != qn('w:val'):
                    del color.attrib[key]
    if s.element.pPr is not None:
        for tag in ['w:pBdr','w:shd']:
            el=s.element.pPr.find(qn(tag))
            if el is not None:
                s.element.pPr.remove(el)
normal=doc.styles['Normal']
normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.space_after=Pt(10)
normal.paragraph_format.widow_control=True
normal.paragraph_format.keep_with_next=False
normal.paragraph_format.keep_together=False
for level,size,before in [(1,14,24),(2,13,10),(3,11,10)]:
    s=doc.styles[f'Heading {level}']
    s.font.size=Pt(size)
    s.font.bold=True
    s.paragraph_format.space_before=Pt(before)
    s.paragraph_format.space_after=Pt(4)
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.keep_together=True
doc.styles['Title'].font.size=Pt(15)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].paragraph_format.line_spacing=1.0
doc.styles['Caption'].paragraph_format.space_after=Pt(8)

def p(text,style='Normal'):
    para=doc.add_paragraph(text,style)
    previous=para._p.getprevious()
    if previous is not None and previous.tag==qn('w:tbl'):
        para.paragraph_format.space_before=Pt(6)
    if style=='Normal':
        para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return para

def h(text,level=1):
    return p(text,f'Heading {level}')

def table(caption,headers,rows,widths):
    cap=p(caption,'Caption')
    cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,width in enumerate(widths):
        t.columns[i].width=Inches(width)
    for cells,values,is_header in [(t.rows[0].cells,headers,True)]+[(t.add_row().cells,row,False) for row in rows]:
        for i,(cell,text) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[i])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(text)
            tcpr=cell._tc.get_or_add_tcPr()
            margin=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'70' if side in ['top','bottom'] else '85');e.set(qn('w:type'),'dxa');margin.append(e)
            tcpr.append(margin)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EAF0' if is_header else 'FFFFFF');tcpr.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0)
                par.paragraph_format.line_spacing=1.0
                par.paragraph_format.keep_with_next=False
                par.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in par.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(9.5);run.font.bold=is_header
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if is_header:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    return t

def fig(filename,caption,full_page=False):
    if full_page:
        doc.add_page_break()
    para=doc.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(0)
    para.paragraph_format.space_after=Pt(0)
    para.paragraph_format.line_spacing=1.0
    para.paragraph_format.keep_with_next=True
    run=para.add_run()
    shape=run.add_picture(str(ROOT/'results/manuscript_figures'/filename),width=Inches(6))
    shape._inline.docPr.set('descr',caption)
    cap=p(caption,'Caption')
    cap.paragraph_format.keep_together=True
    if full_page:
        doc.add_page_break()

title=p('Soil moisture variability and spatial reconstruction in the Aral Sea region','Title')
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in source_author_lines:
    if text.strip():
        a=p(text)
        a.alignment=WD_ALIGN_PARAGRAPH.CENTER
        a.paragraph_format.space_after=Pt(4)
        for r in a.runs:r.font.size=Pt(10)

h('Highlights')
for text in [
    'Two August 2026 snapshots describe soil-water and thermal patterns across the Aral Sea region.',
    'Spatial tests keep both observations at each grid cell together and separate test blocks from training cells.',
    'Gradient boosting gives the lowest pooled soil-moisture error across three spatial validation designs.',
    'Short-interval changes and spatial reconstruction are evaluated as distinct questions.'
]:p(text,'List Bullet')

h('Abstract')
p('The Aral Sea and its surroundings provide a distinct setting for examining spatial relationships between soil water, temperature and atmospheric conditions. This study analyses eight daily ECMWF reanalysis fields for 10 and 18 August 2026 over 43.3–47.5° N and 58.1–63.1° E. The 0.1° grid contains 2,193 locations and 4,386 cell-date records. Area-weighted summaries and paired differences describe conditions across the domain, while six fixed model configurations reconstruct topsoil volumetric water content under buffered spatial validation. Both dates at each location are assigned to the same test block. Mean soil moisture increased from 0.0589 to 0.0607 m³/m³, while mean temperature in the 0–7 cm soil layer decreased by 5.10°C. Under the primary four-quadrant design with a 0.3° buffer, gradient boosting achieved a root mean squared error of 0.0452 m³/m³ and pooled R² of 0.444, compared with 0.0695 m³/m³ for the training-mean baseline. Gradient-boosting errors ranged from 0.0394 to 0.0452 m³/m³ across the three spatial designs. An eight-day persistence benchmark produced an error of 0.0049 m³/m³. The findings describe short-interval changes and the spatial predictability of reanalysis soil moisture within the Aral Sea domain. Spatial block performance and temporal persistence represent different prediction settings and are interpreted separately.')
h('Keywords')
p('Aral Sea; soil moisture; ERA5 reanalysis; spatial validation; gradient boosting; soil temperature')

h('1 Introduction')
p('The contraction of the Aral Sea has transformed the hydrological setting of the lake, its former margins and surrounding land. Irrigation withdrawals from its tributaries reduced inflow and contributed to extensive changes in water area, salinity and exposed lakebed conditions (Micklin, 2007). These changes make the Aral Sea region an important setting for studying environmental contrasts at the land surface. Soil moisture is relevant to that setting because it describes the water held in the near-surface soil layer and varies alongside thermal, vegetation and atmospheric conditions.')
p('Regional gridded data allow those contrasts to be examined on a consistent spatial support. An empirical model can then test whether the soil-water pattern is reconstructable from related environmental fields and geographic position. This objective requires a clear distinction between describing a reanalysis field and validating an environmental process. A model fitted to reanalysis soil moisture evaluates agreement with that modelled state. Its errors also depend on which parts of the domain are withheld during evaluation.')
p('The present study focuses exclusively on the Aral Sea and its surroundings. It examines two daily snapshots, 10 and 18 August 2026, at the archived 0.1° resolution. The objectives are to describe paired changes in soil-water and thermal conditions, compare six spatial reconstruction configurations, and assess sensitivity to the geometry and separation of training and test areas. A separate persistence benchmark describes how well the first soil-moisture field carries forward to the second date. Together, these experiments provide a reproducible account of environmental patterns and model behaviour within one regional domain.')

h('2 Literature review')
for text in review['paragraphs']:
    p(text.replace('surface-water variable','soil-water variable'))

h('3 Study area and data')
h('3.1 Aral Sea domain and observation dates',2)
p('The study domain is a rectangular grid extending from 43.3 to 47.5° N and from 58.1 to 63.1° E, covering the Aral Sea and its surroundings. It contains 43 latitude positions and 51 longitude positions, giving 2,193 spatial locations. Grid spacing is 0.1° in both directions. The two archived daily records are dated 10 and 18 August 2026, separated by eight days. Each variable therefore contributes 4,386 cell-date values. Latitude and longitude coordinates agree exactly across the eight NetCDF files, and all grid records are finite.')
p('The analysis retains the archived grid without spatial interpolation. All regional means and changing-area fractions refer to this complete rectangular domain. They do not represent a shoreline, exposed-seabed or land-cover subset. The NetCDF metadata identify ECMWF as the institution and record daily-mean aggregation with a five-hour time shift. Product naming in this paper is kept at the ERA5 family level, consistent with the archived provenance.')
h('3.2 Variables and physical interpretation',2)
p('Topsoil volumetric water content in the 0–7 cm model layer is the reconstruction target. Environmental predictors comprise soil temperatures at 0–7 and 7–28 cm, the eastward wind component at 10 m, and low-vegetation leaf area index. Latitude and longitude form a separate geographic predictor set. Soil temperatures are converted from kelvin to degrees Celsius. The signed eastward wind component is retained in its original meaning. Low-vegetation leaf area index is interpreted as prescribed climatological vegetation structure rather than a direct observation of recent vegetation change (Muñoz-Sabater et al., 2021; ECMWF, 2026).')
table('Table 1  Reanalysis fields and their roles in the study',
      ['Field','Symbol','Unit','Role'],[
          ['Topsoil volumetric water content','swvl1','m³/m³','Target; 0–7 cm'],
          ['Soil temperature in layer 1','stl1','K, converted to °C','Predictor; 0–7 cm'],
          ['Soil temperature in layer 2','stl2','K, converted to °C','Predictor; 7–28 cm'],
          ['Eastward wind component at 10 m','u10','m/s','Predictor'],
          ['Low-vegetation leaf area index','lai_lv','m²/m²','Climatological predictor'],
          ['Snow cover','snowc','%','Input diagnostic'],
          ['Lake bottom temperature','lblt','K','Input diagnostic'],
          ['Lake ice thickness','licd','m','Input diagnostic'],
      ],[2.5,.65,.9,1.95])
p('Snow cover is zero throughout the two snapshots. Lake fields are retained as input diagnostics. Their interpretation follows the model definition, under which lake quantities can be calculated at grid locations irrespective of actual local water coverage (ECMWF, 2026). Numerical quality control identified 37 soil-water values near −10⁻²⁰ and 4,386 lake-ice values near −10⁻¹⁹. Negative values between −10⁻¹² and zero were treated as numerical roundoff and set to zero for analysis. Original NetCDF values and the raw tabular export were preserved.')

h('4 Methods')
h('4.1 Regional summaries and paired changes',2)
p('Each cell receives a weight proportional to the cosine of its latitude. With equal angular spacing, these weights account for the change in cell area across the domain. Regional means are calculated separately for each date. Paired changes are defined as the value on 18 August minus the value on 10 August at the same grid location. The proportion of the domain showing an increase or decrease is calculated using the same area weights, with changes within ±10⁻¹² treated as unchanged. Medians and spatial correlation coefficients are reported without area weighting. Pearson and Spearman coefficients describe associations pooled across both snapshots.')
h('4.2 Reconstruction models',2)
p('Six fixed configurations are evaluated: an area-weighted training-mean baseline, ridge regression using environmental variables, random forests using geographic variables, environmental variables or both sets, and gradient boosting using both sets. The forest and boosting methods follow the tree-ensemble approaches described by Breiman (2001) and Friedman (2001). Model settings are specified before evaluating the withheld blocks. The ridge scaler is fitted within each training fold. Model fitting and error summaries use cosine-latitude weights, and predictions are retained without clipping or subsequent calibration.')
table('Table 2  Fixed reconstruction configurations',
      ['Configuration','Predictors','Settings'],[
          ['Training mean','None','Area-weighted training target mean'],
          ['Ridge environmental','Temperatures, u10, LAI','Training-fold standardization; penalty 10'],
          ['RF coordinates','Latitude, longitude','200 trees; minimum leaf 10'],
          ['RF environmental','Temperatures, u10, LAI','200 trees; minimum leaf 10'],
          ['RF combined','Environmental and geographic','200 trees; minimum leaf 10'],
          ['GBR combined','Environmental and geographic','150 stages; learning rate 0.03; depth 2; minimum leaf 10'],
      ],[1.35,1.8,2.85])
p('All features are considered at each random-forest split, bootstrap sampling is enabled, and the forests use squared-error splitting. Gradient boosting uses squared-error loss. The random seed is 20261002. Both dates are included in the spatial training pool, so the fitted task is contemporaneous soil-moisture reconstruction at withheld locations. Date is not an input feature. No test score is used to tune a parameter or choose a geographic partition.')
h('4.3 Spatial partitions and evaluation',2)
p('The primary design divides the study rectangle into four quadrants at the midpoint of its latitude and longitude extents. Each quadrant is withheld once, with both observations at a location assigned to the same fold. Training cells within a 0.3° rectangular coordinate buffer around the held-out block are excluded. The distance rule uses the maximum of the latitude and longitude distances to the held rectangle, in degrees. Sensitivity analyses use a 0.6° quadrant buffer and four longitude strips with a 0.3° buffer. These are coordinate buffers, not fixed distances in kilometres. This design responds to the spatial dependence concerns discussed by Roberts et al. (2017).')
p('Each configuration generates one held-out prediction for every cell-date record within each design. Pooled root mean squared error (RMSE), mean absolute error (MAE), bias and R² are recomputed from all held-out predictions using the area weights. Bias is defined as prediction minus reanalysis. Pooled R² uses the weighted mean of all evaluation targets as its reference, whereas fold-specific R² uses the target mean within that fold. MSE skill expresses one minus the ratio of model mean squared error to that of the fold-specific training-mean baseline. Fold ranges describe variation among the four spatial blocks; they are not confidence intervals.')
h('4.4 Persistence and reproducibility',2)
p('A separate eight-day persistence benchmark uses the 10 August soil-moisture value at each location as its prediction for 18 August. This benchmark uses the earlier target field directly, whereas spatial reconstruction estimates targets at locations excluded from model fitting. Their errors therefore answer different questions and are not ranked against one another. The persistence calculation evaluates one temporal pair.')
p('The analysis is implemented in Python with NumPy 2.5.3, pandas 3.0.6, xarray 2026.9.0, netCDF4 1.7.4, SciPy 1.18.1 and scikit-learn 1.9.1. The three spatial designs, six configurations and four folds produce 72 evaluations and 78,948 held-out predictions. Independent calculations verify prediction coverage, geographic separation, error metrics, paired changes and input-file hashes. All 20 verification checks pass.')

h('5 Results')
h('5.1 Soil-water and thermal changes',2)
p('Mean topsoil volumetric water content increased from 0.0589 to 0.0607 m³/m³, an area-weighted difference of 0.0018 m³/m³ over eight days (Table 3). Soil moisture increased over 54.97% of the grid-domain area and decreased over 33.28%; the remaining 11.75% was unchanged at the numerical tolerance used. The median cell-level difference was 0.0000725 m³/m³, smaller than the weighted mean change. The paired maps show that moisture gains are spatially concentrated, particularly toward the southern part of the domain (Figure 1).')
table('Table 3  Area-weighted regional conditions on the two dates',
      ['Variable','Unit','10 August','18 August','Difference'],[
          ['Topsoil water','m³/m³','0.0589','0.0607','+0.0018'],
          ['Soil temperature 0–7 cm','°C','34.85','29.75','−5.10'],
          ['Soil temperature 7–28 cm','°C','32.15','28.77','−3.38'],
          ['Eastward wind at 10 m','m/s','−1.65','−2.52','−0.86'],
          ['Low-vegetation LAI','m²/m²','0.4550','0.4533','−0.0017'],
      ],[2.05,.65,1.05,1.05,1.2])
p('Temperature decreased at every grid location in both soil layers. Mean temperature fell by 5.10°C in the 0–7 cm layer and by 3.38°C in the 7–28 cm layer. The eastward wind component became more negative, changing from −1.65 to −2.52 m/s. The small LAI difference is interpreted within its prescribed seasonal climatology. These contrasts describe the conditions represented on the two August dates.')
fig('figure_1_moisture_maps.png','Figure 1. Topsoil volumetric water content on 10 and 18 August 2026 and the paired difference. The first two maps use the same scale; the difference is 18 August minus 10 August. Values retain the 0.1° grid and represent the complete study rectangle.')

h('5.2 Environmental associations',2)
p('Across the 4,386 cell-date records, soil moisture had a Spearman correlation of 0.716 with low-vegetation LAI and 0.545 with the eastward wind component. Corresponding Pearson correlations were 0.580 and 0.497. Associations with soil temperature were weak in the pooled sample: Spearman coefficients were −0.043 for the upper layer and −0.038 for the deeper layer. The two temperature variables were strongly associated with one another, with a Spearman coefficient of approximately 0.97. These are descriptive associations within the reanalysis fields; repeated locations and spatial dependence are retained in their interpretation.')
h('5.3 Performance under the primary spatial design',2)
p('Combined gradient boosting achieved the lowest primary-design RMSE, 0.0452 m³/m³, with MAE 0.0318 m³/m³ and pooled R² 0.444 (Table 4). Its RMSE was 34.9% lower than the training-mean baseline, corresponding to 57.7% MSE skill. Ridge regression achieved RMSE 0.0488 m³/m³, followed by the environmental random forest at 0.0521 m³/m³. The combined forest had RMSE 0.0540 m³/m³, while the coordinate-only forest had RMSE 0.0571 m³/m³. Thus, adding location to the forest did not improve its pooled primary-design error relative to the environmental-only forest.')
model_names={'training_mean':'Training mean','ridge_environment':'Ridge environmental','rf_geography':'RF coordinates','rf_environment':'RF environmental','rf_combined':'RF combined','gbr_combined':'GBR combined'}
order=list(model_names)
rows=[]
for name in order:
    x=metrics[('quadrants_buffer_0.3',name)]
    rows.append([model_names[name],f"{x['rmse']:.4f}",f"{x['mae']:.4f}",f"{x['bias']:.4f}",f"{x['r2']:.3f}",f"{100*x['mse_skill_vs_training_mean']:.1f}"])
table('Table 4  Pooled performance for quadrants with a 0.3 degree buffer',
      ['Configuration','RMSE','MAE','Bias','R²','MSE skill %'],rows,[1.7,.86,.86,.86,.72,1.0])
p('RMSE, MAE and bias are in m³/m³. Values use all held-out cell-date predictions and cosine-latitude weights. The training-mean reference is recalculated in each fold. RF denotes random forest and GBR denotes gradient-boosting regression.','Caption')
p('Spatial performance varied among held-out blocks. Gradient-boosting RMSE ranged from 0.0125 to 0.0549 m³/m³, and fold-specific R² ranged from −0.337 to 0.489. The negative value in the southwestern block shows that pooled improvement does not imply equally strong reconstruction throughout the region. The combined forest was also heterogeneous, with fold RMSE from 0.0120 to 0.0682 m³/m³. Consequently, regional pooled scores are read alongside block-specific behaviour.')
h('5.4 Sensitivity to spatial separation',2)
p('Gradient boosting retained the lowest pooled RMSE among the six configurations under both alternative partitions (Figure 2). Its RMSE was 0.0431 m³/m³ with a 0.6° quadrant buffer and 0.0394 m³/m³ with 0.3°-buffered longitude strips; corresponding R² values were 0.494 and 0.577. The combined forest reached RMSE 0.0435 and 0.0399 m³/m³ in these two settings. Increasing the buffer therefore did not produce a monotonic increase in pooled error. The result reflects changes in the composition and distribution of retained training cells, as well as separation from the test block.')
fig('figure_2_model_comparison.png','Figure 2. Pooled RMSE for all six fixed configurations across the primary quadrant design and two spatial sensitivity designs. All points refer to contemporaneous reconstruction of reanalysis soil moisture; both dates at each location remain in the same held-out block.')
h('5.5 Eight-day persistence',2)
p('Carrying the 10 August moisture field forward to 18 August produced RMSE 0.0049 m³/m³, MAE 0.0026 m³/m³ and bias −0.0018 m³/m³. The negative bias follows from the increase in the regional mean between dates. R² was 0.993, reflecting the persistence of the broad spatial moisture pattern over this interval. Because the benchmark has access to the previous soil-moisture value at each evaluation location, it is a distinct temporal reference rather than a competing spatial reconstruction model.')

h('6 Discussion')
h('6.1 Interpreting the Aral Sea patterns',2)
p('The two snapshots combine a regional cooling signal with a small mean increase in topsoil water. The difference between the median moisture change and the weighted mean indicates that the regional mean is influenced by spatially concentrated changes. This supports examining maps and distributions together when describing the Aral Sea environment. The observed eight-day contrast is a short-interval land-state comparison. It does not measure the rate of the historical desiccation described by Micklin (2007).')
p('The contrast between stronger associations with LAI and wind and weaker pooled associations with temperature also requires care. Temperature differs substantially between dates, while much of the spatial moisture pattern persists. A correlation calculated across both dates mixes geographic variation with date-to-date variation. Climatological vegetation structure and atmospheric fields are useful descriptive covariates in this reconstruction setting, but their correlations do not establish which process caused a local moisture change. This interpretation is consistent with treating the reanalysis as a coupled source of modelled environmental fields (Muñoz-Sabater et al., 2021).')
h('6.2 What spatial validation shows',2)
p('The consistent pooled advantage of gradient boosting across the three predefined designs supports its use as a compact empirical description of the regional reanalysis pattern under the tested conditions. It does not establish a universally optimal model. Ridge regression outperformed the forests in the primary design, and the benefit of adding geographic coordinates to a forest changed with the partition. These results show why model complexity and larger predictor sets should be assessed through explicit comparisons rather than assumed to improve reconstruction.')
p('The block-level errors are particularly relevant in a heterogeneous domain. A negative southwestern-block R² for gradient boosting coexists with a positive pooled R² because the pooled and local reference variances differ and performance varies geographically. Likewise, improved error after widening a buffer is possible when the remaining training distribution changes. The sensitivity results therefore describe dependence on partition geometry and training composition; they should not be interpreted as a controlled experiment isolating distance alone. Reporting all predefined designs follows the validation principles outlined by Roberts et al. (2017).')
h('6.3 Temporal persistence and scientific scope',2)
p('The low persistence error shows that the spatial moisture field changes modestly relative to its between-location variation over the eight-day interval. Its high R² is compatible with locally meaningful moisture gains because the dominant regional contrasts remain present. The benchmark is useful as a reference for this temporal pair, while the learned models answer the separate question of reconstruction at spatially withheld sites. Keeping these two tasks distinct avoids attributing temporal forecast skill to models trained using both dates.')
p('The conclusions apply to modelled soil-water values on the 0.1° Aral Sea grid and to the two specified August dates. The analysis describes internal spatial consistency and paired land-state changes, rather than direct measurements of salinity or a multiyear environmental trend. Within this scope, the reproducible workflow and the reported block-level variability provide a clear basis for comparing regional soil-moisture reconstruction methods.')

h('7 Conclusions')
p('This study provides an Aral Sea-focused analysis of two daily reanalysis snapshots and a spatial comparison of six soil-moisture reconstruction configurations. Between 10 and 18 August 2026, regional mean topsoil water increased by 0.0018 m³/m³ while upper-layer soil temperature decreased by 5.10°C. Moisture changes were spatially uneven, and the broader regional pattern persisted across the interval.')
p('Gradient boosting achieved the lowest pooled error in each predefined spatial design, with RMSE from 0.0394 to 0.0452 m³/m³. Its primary-design pooled R² of 0.444 was accompanied by substantial variation between geographic blocks. These findings establish the importance of reporting split geometry, local errors and simple reference models when assessing soil-moisture reconstruction in the Aral Sea and its surroundings. The conclusions concern the analysed grid and dates, with temporal persistence and spatial prediction evaluated separately.')

h('Declaration of competing interest')
p('The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.')
h('Data and code availability')
p('The study archive contains the eight input NetCDF files, their SHA-256 identifiers, derived regional tables, paired changes, spatial fold definitions and all held-out predictions. Reproducible Python scripts and package versions accompany the analysis. The underlying ECMWF climate products are distributed through the Copernicus Climate Data Store.')
h('References')
refs=[
    'Breiman, L. (2001). Random forests. Machine Learning, 45, 5–32. https://doi.org/10.1023/A:1010933404324',
    'Copernicus Climate Change Service (C3S) (2019). ERA5-Land hourly data from 1950 to present. Climate Data Store. https://doi.org/10.24381/cds.e2161bac (accessed 2 October 2026).',
    'European Centre for Medium-Range Weather Forecasts (ECMWF) (2026). ERA5-Land data documentation. https://confluence.ecmwf.int/spaces/CKB/pages/140385202/ERA5-Land+data+documentation (accessed 2 October 2026).',
    'Friedman, J. H. (2001). Greedy function approximation: A gradient boosting machine. The Annals of Statistics, 29(5), 1189–1232. https://doi.org/10.1214/aos/1013203451',
    'Micklin, P. (2007). The Aral Sea disaster. Annual Review of Earth and Planetary Sciences, 35, 47–72. https://doi.org/10.1146/annurev.earth.35.031306.140120',
    'Muñoz-Sabater, J., Dutra, E., Agustí-Panareda, A., et al. (2021). ERA5-Land: a state-of-the-art global reanalysis dataset for land applications. Earth System Science Data, 13, 4349–4383. https://doi.org/10.5194/essd-13-4349-2021',
    'Roberts, D. R., Bahn, V., Ciuti, S., et al. (2017). Cross-validation strategies for data with temporal, spatial, hierarchical, or phylogenetic structure. Ecography, 40(8), 913–929. https://doi.org/10.1111/ecog.02881',
    'Small, E. E., Sloan, L. C., Hostetler, S., and Giorgi, F. (1999). Simulating the water balance of the Aral Sea with a coupled regional climate-lake model. Journal of Geophysical Research: Atmospheres, 104(D6), 6583–6602. https://doi.org/10.1029/98JD02348',
]
for ref in refs:
    a=p(ref)
    a.alignment=WD_ALIGN_PARAGRAPH.LEFT
    a.paragraph_format.space_after=Pt(7)
    a.paragraph_format.keep_together=True
    a.paragraph_format.left_indent=Inches(.2)
    a.paragraph_format.first_line_indent=Inches(-.2)
    for run in a.runs:run.font.size=Pt(10)

doc.core_properties.title='Soil moisture variability and spatial reconstruction in the Aral Sea region'
doc.core_properties.subject='Aral Sea regional reanalysis and spatial soil-moisture reconstruction'
doc.core_properties.keywords='Aral Sea, soil moisture, ERA5, spatial validation'
doc.core_properties.comments=''
doc.save(OUTPUT)

# Record the authored text and revision scope for independent checking.
reopened=Document(OUTPUT)
text='\n\n'.join(p.text for p in reopened.paragraphs)
(QA/'revised_text.txt').write_text(text,encoding='utf-8')
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
assert all(term not in text for term in ['Bozaigyr','Atyrau','Shol','SDRI','CSI','missing data','unavailable data'])
record={
    'source':str(SOURCE),'source_sha256':source_hash,'output':str(OUTPUT),
    'output_sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    'scope':'Aral Sea and surroundings; two-date reanalysis spatial reconstruction',
    'dates':summary['analysis_dates'],'paragraphs':len(reopened.paragraphs),
    'tables':len(reopened.tables),'figures':len(reopened.inline_shapes),
    'changes':['Rewrote literature review and discussion around Aral Sea sources and verified results',
               'Replaced other-region and unperformed fusion modules with the completed experiments',
               'Removed missing-data inventories and retained the scientific scope in methods and conclusions',
               'Preserved source page geometry, typography and author metadata'],
    'source_unchanged':True,
}
(QA/'revision_manifest.json').write_text(json.dumps(record,indent=2),encoding='utf-8')
print(json.dumps(record,indent=2))
